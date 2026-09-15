# -*- coding: utf-8 -*-
# ============================================================================
# excel_to_ics.py
# ---------------
# Конвертує розклад з Excel-файлу у формат iCalendar (.ics), який можна
# імпортувати в Google Calendar.
#
# Структура вхідного Excel-файлу (приклад "11.xlsx"):
#   - Рядок із днем тижня ("Пн", "Вт", ...) є "маркером дня". У наступних
#     колонках цього рядка зберігаються дати (напр. 01.09.2026, 08.09.2026...).
#   - Під кожним маркером дня йдуть рядки "пари", напр. "1 пара 09:00-10:20",
#     де в першій колонці вказано номер і час пари, а в інших колонках —
#     предмети на конкретні дати.
#
# Приклад рядка пари:
#     "1 пара 09:00-10:20" | "ВМ [Пз]\nауд. 1-426_[6]\nПМ-26" | ...
#   Текст осередку розділено на рядки:
#     - рядок 0 (перший)     -> назва предмета (SUMMARY)
#     - рядок 1 (другой)     -> аудиторія (LOCATION)
#     - рядок 2 і далі       -> додатковий опис (DESCRIPTION)
#
# Програма обходить усі пари по всіх датах, створює окрему подію для кожного
# поєднання "дата + пара", і записує результат в .ics файл.
# ============================================================================

import sys                      # для роботи з аргументами командного рядка та виходу
import re                       # регулярні вирази — для пошуку часу в тексті пари
import uuid                     # генерація унікальних ідентифікаторів (UID) подій
from pathlib import Path        # зручна робота зі шляхами до файлів
from datetime import datetime   # робота з датами і часом

import openpyxl                 # бібліотека для читання/запису Excel-файлів


# Множина назв днів тижня (українською), що зустрічаються як маркери дня в Excel.
DAY_NAMES = {"Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Нд"}


def escape_ics(text):
    """
    Екранує спеціальні символи відповідно до специфікації iCalendar.

    У форматі ICS символи \\, ;  та , мають особливе значення, тому їх
    потрібно "екранувати" (попереджувати зворотньою косою). Символ нового
    рядка \\n теж замінюється на текстову послідовність "\\n", щоб запис
    залишився одним рядком у файлі.

    Аргументи:
        text: рядок, який треба екранувати (або None).

    Повертає:
        Екранований рядок. Якщо text є None — порожній рядок.
    """
    if text is None:               # якщо значення відсутнє — повертаємо порожній рядок
        return ""

    text = str(text)               # гарантуємо, що працюємо саме з рядком

    # Заміна спеціальних символів ICS на екрановані послідовності.
    # Порядок має значення: зворотню косу замінюємо першою.
    text = text.replace("\\", "\\\\")   #  \  ->  \\
    text = text.replace(";", r"\;")     #  ;  ->  \;
    text = text.replace(",", r"\,")     #  ,  ->  \,
    text = text.replace("\n", r"\n")    #  новий рядок -> "\n"

    return text


def parse_time_range(cell_value):
    """
    Шукає в тексті пари часовий проміжок виду "ГГ:ХХ-ГГ:ХХ".

    Наприклад, для "1 пара 09:00-10:20" поверне кортеж ("09:00", "10:20").

    Аргументи:
        cell_value: вміст клітинки (може бути рядком або іншим типом).

    Повертає:
        Кортеж (start, end) або None, якщо проміжок знайти не вдалося.
    """
    # Регулярний вираз: дві цифри ":" дві цифри, дефіс, і ще раз той самий шаблон.
    m = re.search(r"(\d{2}:\d{2})-(\d{2}:\d{2})", str(cell_value))

    if not m:                      # якщо збігів немає — час не знайдено
        return None

    # group(1) — час початку, group(2) — час кінця.
    return m.group(1), m.group(2)


def create_event(date_str, time_start, time_end, cell_text):
    """
    Створює словник з даними однієї події на основі вмісту клітинки.

    Вміст клітинки з предметом зазвичай містить кілька рядків:
        рядок 0 -> назва предмета (SUMMARY)
        рядок 1 -> аудиторія      (LOCATION)
        рядок 2+ -> опис          (DESCRIPTION)

    Аргументи:
        date_str:   дата у вигляді рядка (або об'єкта datetime).
        time_start: час початку, "ГГ:ХХ".
        time_end:   час кінця, "ГГ:ХХ".
        cell_text:  текст клітинки з інформацією про предмет.

    Повертає:
        Словник з полями uid, summary, location, description, start, end.
    """
    # Розбиваємо вміст клітинки на окремі рядки.
    parts = str(cell_text).split("\n")

    # Перший рядок — назва предмета (або порожній рядок, якщо рядків немає).
    subject = parts[0].strip() if len(parts) > 0 else ""

    # Другий рядок — аудиторія (якщо існує).
    location = parts[1].strip() if len(parts) > 1 else ""

    # Решта рядків — опис, об'єднуємо назад через перехід на новий рядок.
    description = "\n".join(p.strip() for p in parts[2:]) if len(parts) > 2 else ""

    # Обробка дати: вона може бути передана як рядок або як об'єкт datetime.
    # Якщо це datetime — перетворюємо у зручний формат ДД.ММ.РРРР.
    if isinstance(date_str, datetime):
        date_str = date_str.strftime("%d.%m.%Y")

    # Формуємо повні дати початку та кінця: "ДД.ММ.РРРР ГГ:ХХ".
    start_dt = datetime.strptime(f"{date_str} {time_start}", "%d.%m.%Y %H:%M")
    end_dt = datetime.strptime(f"{date_str} {time_end}", "%d.%m.%Y %H:%M")

    # Повертаємо словник з усіма даними події.
    return {
        "uid": str(uuid.uuid4()),          # унікальний ідентифікатор події
        "summary": subject,                # назва предмета
        "location": location,              # аудиторія
        "description": description,        # опис
        "start": start_dt,                 # дата+час початку
        "end": end_dt,                     # дата+час кінця
    }


def extract_events(ws):
    """
    Обходить усі рядки робочого аркуша Excel і збирає всі події.

    Логіка обходу:
      1. Зустрічаємо рядок-маркер дня (напр. "Пн") — запам'ятовуємо цей рядок
         як "date_row". У ньому в колонках 2..N розташовані дати.
      2. Зустрічаємо рядок пари (напр. "1 пара 09:00-10:20") — отримуємо час.
         Далі йдемо вниз по рядках (scan_row), поки не зустрінемо новий маркер
         (день, пара чи час). Для кожного рядка перевіряємо всі колонки з
         датами і, якщо є предмет, створюємо подію.

    Аргументи:
        ws: активний робочий аркуш openpyxl.

    Повертає:
        Список словників-подій (результат функції create_event).
    """
    events = []          # тут накопичуємо всі знайдені події
    date_row = None      # рядок маркера дня (в ньому дати)
    row_idx = 1          # починаємо з першого рядка (1=перший рядок Excel)

    # Цикл по всіх рядках аркуша.
    while row_idx <= ws.max_row:
        # Вміст першої клітинки поточного рядка.
        first_cell = ws.cell(row_idx, 1).value

        # Якщо це маркер дня — запам'ятовуємо номер рядка і йдемо далі.
        if isinstance(first_cell, str) and first_cell.strip() in DAY_NAMES:
            date_row = row_idx          # запам'ятовуємо рядок з датами
            row_idx += 1
            continue                    # переходимо до наступного рядка

        # Якщо це рядок пари (містить слово "пара").
        if isinstance(first_cell, str) and "пара" in first_cell.lower():
            # Витягуємо часовий проміжок з тексту пари.
            time_range = parse_time_range(first_cell)

            if not time_range:          # якщо час не знайдено — пропускаємо рядок
                row_idx += 1
                continue

            time_start, time_end = time_range   # час початку та кінця пари
            scan_row = row_idx                  # починаємо сканувати з цього ж рядка

            # Скануємо вниз, доки не натрапимо на новий маркер (день/пара/час).
            while scan_row <= ws.max_row:
                # Для рядків НИЖЧЕ поточного перевіряємо, чи це новий маркер —
                # якщо так, припиняємо сканування поточної пари.
                if scan_row > row_idx:
                    next_val = ws.cell(scan_row, 1).value
                    if isinstance(next_val, str):
                        if (next_val.strip() in DAY_NAMES or      # новий день
                            "пара" in next_val.lower() or         # нова пара
                            parse_time_range(next_val)):          # або новий час
                            break                                  # стоп — кінець пари

                # Якщо вже є рядок дня з датами — обробляємо колонки.
                if date_row is not None:
                    # Перебираємо всі колонки, починаючи з другої (дати).
                    for col in range(2, ws.max_column + 1):
                        date_value = ws.cell(date_row, col).value     # дата
                        lesson_value = ws.cell(scan_row, col).value   # предмет

                        # Пропускаємо, якщо дати або предмета немає.
                        if not date_value or not lesson_value:
                            continue

                        # Перетворюємо дату в рядок (якщо це datetime — форматуємо).
                        if isinstance(date_value, datetime):
                            date_str = date_value.strftime("%d.%m.%Y")
                        else:
                            date_str = str(date_value).strip()

                        # Створюємо подію; якщо виникає помилка — повідомляємо,
                        # але не перериваємо обробку решти файлу.
                        try:
                            event = create_event(date_str, time_start, time_end, lesson_value)
                            events.append(event)   # додаємо подію до списку
                            print(f"[row {scan_row}] {date_str} {time_start}-{time_end} → {lesson_value}")
                        except Exception as e:
                            print(f"Ошибка при создании события: {e}")

                scan_row += 1            # переходимо до наступного рядка сканування

        row_idx += 1                     # переходимо до наступного рядка загального циклу

    return events                        # повертаємо всі знайдені події


def save_ics(events, output_file):
    """
    Записує список подій у файл формату iCalendar (.ics).

    Формат ICS — це текстовий файл з ієрархією блоків:
      BEGIN:VCALENDAR
        BEGIN:VEVENT ... END:VEVENT   (одна подія на заняття)
        ...
      END:VCALENDAR

    Аргументи:
        events:      список словників-подій (як з extract_events).
        output_file: шлях до вихідного .ics файлу.
    """
    # Заголовки календаря (обов'язкові та службові записи ICS).
    lines = [
        "BEGIN:VCALENDAR",                       # початок календаря
        "VERSION:2.0",                           # версія формату
        "PRODID:-//ExcelToICS//UA//",            # ідентифікатор виробника
        "CALSCALE:GREGORIAN",                    # календар (григоріанський)
        "METHOD:PUBLISH",                        # метод публікації
        "X-WR-TIMEZONE:Europe/Kyiv",             # часовий пояс (Київ)
    ]

    # Поточний час у форматі ICS (UTC) — використовується як мітка створення.
    now = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")

    # Для кожної події формуємо блок VEVENT з усіма її даними.
    for event in events:
        lines.extend([
            "BEGIN:VEVENT",                                                          # початок події
            f"UID:{event['uid']}",                                                  # унікальний ID
            f"DTSTAMP:{now}",                                                       # час створення запису
            f"DTSTART:{event['start'].strftime('%Y%m%dT%H%M%S')}",                  # дата/час початку
            f"DTEND:{event['end'].strftime('%Y%m%dT%H%M%S')}",                      # дата/час кінця
            f"SUMMARY:{escape_ics(event['summary'])}",                              # назва предмета (екрановано)
            f"LOCATION:{escape_ics(event['location'])}",                            # аудиторія (екрановано)
            f"DESCRIPTION:{escape_ics(event['description'])}",                      # опис (екрановано)
            "END:VEVENT",                                                           # кінець події
        ])

    lines.append("END:VCALENDAR")               # кінець календаря

    # Записуємо весь вміст у файл (кожен запис — окремий рядок), в кодуванні UTF-8.
    Path(output_file).write_text("\n".join(lines), encoding="utf-8")
    print(f"ICS файл сохранён: {output_file}")


# Головний блок: виконується лише коли файл запускають напряму (не імпортують).
if __name__ == "__main__":
    # Перевіряємо, що передано шлях до Excel-файлу.
    if len(sys.argv) < 2:
        print("Использование: python excel_to_ics.py <файл.xlsx>")
        sys.exit(1)                       # завершуємо з кодом помилки

    input_file = sys.argv[1]                                 # шлях вхідного файлу
    output_file = Path(input_file).stem + ".ics"             # ім'я виводу = ім'я вхідного, розширення .ics

    wb = openpyxl.load_workbook(input_file)                  # відкриваємо книгу Excel
    ws = wb.active                                           # беремо активний аркуш

    events = extract_events(ws)                              # витягуємо всі події
    save_ics(events, output_file)                            # записуємо в .ics файл

    # Повідомляємо користувачу кількість імпортованих подій.
    print(f"Всього подій: {len(events)}")
    print(f"Знайдено {len(events)} подій у файлі '{input_file}' та збережено до '{output_file}'")
    sys.exit(0)                              # успішне завершення
